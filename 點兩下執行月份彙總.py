#!/usr/bin/env python3
"""
讀取「支出.xlsx」的各明細工作表，彙整成「月份彙總.xlsx」。

輸出檔內容：
  1. 月份彙總  ：每個月一列（1月～12月），每個產品兩欄（數量、金額），
                 最右邊是當月合計金額，最下面一列是全年總結。
  2. 月份明細  ：長表格式（月份 / 產品名稱 / 數量 / 金額），
                 每月底下有小計，最後有總結。

資料夾結構（年份資料夾用民國年命名，放在本程式旁邊）：
  中油進貨表/
    點兩下執行月份彙總.py
    114/支出.xlsx  →  114/月份彙總.xlsx
    115/支出.xlsx  →  115/月份彙總.xlsx

用法（執行前會自動 git pull 拉最新版；直接點兩下執行等同不帶參數）：
  python3 點兩下執行月份彙總.py                      # 處理今年與前一年（例如 115、114）
  python3 點兩下執行月份彙總.py --push               # 產出後自動 commit + push 月份彙總.xlsx
  python3 點兩下執行月份彙總.py --no-pull            # 不要先 git pull
  python3 點兩下執行月份彙總.py 113                  # 只處理指定年份（民國或西元皆可，可多個）
  python3 點兩下執行月份彙總.py --base /其他/目錄     # 年份資料夾放在別的地方
  python3 點兩下執行月份彙總.py --keep-month-suffix  # 不合併「勞保費-1月」「勞保費-2月」這類名稱
"""
import argparse
import datetime as dt
import re
import subprocess
from pathlib import Path
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONTHS = list(range(1, 13))
HEADER_DATE = "日期"
HEADER_NAME = "商品名稱"
HEADER_QTY = "數量"
HEADER_AMT = "金額"

# 輸出時排在最前面的產品（依此順序）；其餘產品依在來源檔中首次出現的順序排在後面
PRIORITY_PRODUCTS = [
    "92汽油", "92折讓",
    "95汽油", "95折讓",
    "柴油", "柴油折讓", "柴油折讓-大宗",
    "尿素", "尿素折讓",
]


# ----------------------------------------------------------------------
# 讀取
# ----------------------------------------------------------------------
def find_header(ws, max_scan=10):
    """在前幾列中找出含有「日期」與「商品名稱」的標題列，回傳 (列號, {欄名: 欄索引})。"""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        values = [c.value for c in ws[r]]
        texts = [str(v).strip() if v is not None else "" for v in values]
        if HEADER_DATE in texts and HEADER_NAME in texts:
            cols = {t: i for i, t in enumerate(texts) if t}
            return r, cols
    return None, None


def parse_month(value, year_filter=None):
    """把日期欄位轉成月份 (1~12)。無法判讀時回傳 None。"""
    if isinstance(value, dt.datetime) or isinstance(value, dt.date):
        if year_filter and value.year != year_filter:
            return "SKIP"
        return value.month
    if isinstance(value, str):
        # 例如 '2/30'、'2026/2/30'、'02-15'
        m = re.match(r"^\s*(?:(\d{4})[/\-.])?(\d{1,2})[/\-.](\d{1,2})\s*$", value)
        if m:
            year, month = m.group(1), int(m.group(2))
            if year_filter and year and int(year) != year_filter:
                return "SKIP"
            if 1 <= month <= 12:
                return month
    return None


def to_number(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return 0


CN_MONTH = r"(?:十[一二]?|[一二三四五六七八九])月"
TRAILING_MONTH = re.compile(r"[\s\-－_]*(?:1[0-2]|0?[1-9]|" + CN_MONTH[:-1] + r")\s*月\s*$")
LEADING_MONTH = re.compile(r"^" + CN_MONTH + r"\s*")


def normalize_name(name, strip_month=True):
    """整理產品名稱：壓縮空白；預設把「-1月」「12月」「一月」這類月份字樣去掉，
    讓「勞保費-1月」「勞保費-2月」合併成同一個產品（報表本身已按月份分列）。"""
    name = re.sub(r"\s+", " ", str(name)).strip()
    if strip_month:
        stripped = TRAILING_MONTH.sub("", name)
        stripped = LEADING_MONTH.sub("", stripped)
        if stripped:
            name = stripped.strip()
    return name


def read_expenses(path, year_filter=None, strip_month=True):
    """
    回傳：
      data[產品][月份] = {"qty": 數量, "amt": 金額}
      products: 依首次出現順序排列的產品名稱清單
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    data = defaultdict(lambda: defaultdict(lambda: {"qty": 0, "amt": 0}))
    products = OrderedDict()
    skipped = []

    for ws in wb.worksheets:
        header_row, cols = find_header(ws)
        if header_row is None:
            continue  # 例如「數量」庫存表，沒有明細欄位就略過

        c_date, c_name = cols[HEADER_DATE], cols[HEADER_NAME]
        c_qty, c_amt = cols.get(HEADER_QTY), cols.get(HEADER_AMT)
        current_month = None

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            raw_date = row[c_date] if c_date < len(row) else None
            name = row[c_name] if c_name < len(row) else None

            if raw_date is not None:
                month = parse_month(raw_date, year_filter)
                if month == "SKIP":
                    current_month = "SKIP"
                    continue
                if month is not None:
                    current_month = month
            if current_month == "SKIP":
                continue

            if name is None or str(name).strip() == "":
                continue  # 空白列 / 小計列
            if current_month is None:
                skipped.append((ws.title, name))
                continue

            name = normalize_name(name, strip_month)
            qty = to_number(row[c_qty]) if c_qty is not None and c_qty < len(row) else 0
            amt = to_number(row[c_amt]) if c_amt is not None and c_amt < len(row) else 0

            products.setdefault(name, None)
            data[name][current_month]["qty"] += qty
            data[name][current_month]["amt"] += amt

    if skipped:
        print(f"⚠ 有 {len(skipped)} 列找不到日期，已略過：", skipped[:5], "...")
    return data, order_products(list(products.keys()))


def order_products(products):
    """PRIORITY_PRODUCTS 裡有的排最前面（依清單順序），其餘維持首次出現順序。"""
    first = [n for n in PRIORITY_PRODUCTS if n in products]
    rest = [n for n in products if n not in PRIORITY_PRODUCTS]
    return first + rest


# ----------------------------------------------------------------------
# 樣式
# ----------------------------------------------------------------------
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
QTY_FMT = "#,##0"
AMT_FMT = "#,##0_);[Red]\\(#,##0\\)"  # 與來源「金額」「總金額」欄相同的整數格式


def style_range(ws, min_row, max_row, min_col, max_col, fill=None, font=None, align=None):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = BORDER
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if align:
                cell.alignment = align


# ----------------------------------------------------------------------
# 工作表 1：每月一列（寬表）
# ----------------------------------------------------------------------
def write_wide_sheet(wb, data, products):
    ws = wb.active
    ws.title = "月份彙總"

    # 標題列（兩列）：第一列產品名稱（合併兩欄），第二列 數量 / 金額
    ws.cell(row=1, column=1, value="月份")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    for i, name in enumerate(products):
        col = 2 + i * 2
        ws.cell(row=1, column=col, value=name)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        ws.cell(row=2, column=col, value="數量")
        ws.cell(row=2, column=col + 1, value="金額")
    total_col = 2 + len(products) * 2
    ws.cell(row=1, column=total_col, value="當月合計金額")
    ws.merge_cells(start_row=1, start_column=total_col, end_row=2, end_column=total_col)

    # 資料列：1月～12月，沒有資料的月份留空
    first_data_row = 3
    for m_idx, month in enumerate(MONTHS):
        r = first_data_row + m_idx
        ws.cell(row=r, column=1, value=f"{month}月")
        amt_cells = []
        for i, name in enumerate(products):
            col = 2 + i * 2
            rec = data[name].get(month)
            if rec is not None:
                ws.cell(row=r, column=col, value=rec["qty"]).number_format = QTY_FMT
                ws.cell(row=r, column=col + 1, value=rec["amt"]).number_format = AMT_FMT
            else:
                ws.cell(row=r, column=col).number_format = QTY_FMT
                ws.cell(row=r, column=col + 1).number_format = AMT_FMT
            amt_cells.append(f"{get_column_letter(col + 1)}{r}")
        has_data = any(month in data[name] for name in products)
        if has_data:
            # 只加總金額欄（隔一欄取一個）
            formula = "=" + "+".join(amt_cells) if amt_cells else "=0"
            ws.cell(row=r, column=total_col, value=formula).number_format = AMT_FMT
        else:
            ws.cell(row=r, column=total_col).number_format = AMT_FMT

    last_data_row = first_data_row + len(MONTHS) - 1

    # 總結列
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="總結")
    for col in range(2, total_col + 1):
        letter = get_column_letter(col)
        c = ws.cell(row=total_row, column=col,
                    value=f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
        c.number_format = AMT_FMT if (col == total_col or (col - 2) % 2 == 1) else QTY_FMT

    # 樣式
    style_range(ws, 1, 2, 1, total_col, fill=HEAD_FILL, font=BOLD, align=CENTER)
    style_range(ws, first_data_row, last_data_row, 1, total_col)
    style_range(ws, total_row, total_row, 1, total_col, fill=TOTAL_FILL, font=BOLD)
    for r in range(first_data_row, total_row + 1):
        ws.cell(row=r, column=1).alignment = CENTER
    ws.column_dimensions["A"].width = 8
    for i in range(len(products)):
        ws.column_dimensions[get_column_letter(2 + i * 2)].width = 10
        ws.column_dimensions[get_column_letter(3 + i * 2)].width = 14
    ws.column_dimensions[get_column_letter(total_col)].width = 16
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B3"
    return ws


# ----------------------------------------------------------------------
# 工作表 2：長表（月份 / 產品 / 數量 / 金額）
# ----------------------------------------------------------------------
def write_long_sheet(wb, data, products):
    ws = wb.create_sheet("月份明細")
    headers = ["月份", "產品名稱", "數量", "金額"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_range(ws, 1, 1, 1, 4, fill=HEAD_FILL, font=BOLD, align=CENTER)

    r = 2
    subtotal_rows = []
    for month in MONTHS:
        rows_this_month = [(name, data[name][month]) for name in products if month in data[name]]
        start = r
        if rows_this_month:
            for name, rec in rows_this_month:
                ws.cell(row=r, column=1, value=f"{month}月")
                ws.cell(row=r, column=2, value=name)
                ws.cell(row=r, column=3, value=rec["qty"]).number_format = QTY_FMT
                ws.cell(row=r, column=4, value=rec["amt"]).number_format = AMT_FMT
                r += 1
            style_range(ws, start, r - 1, 1, 4)
        else:
            ws.cell(row=r, column=1, value=f"{month}月")  # 沒有資料的月份留空
            style_range(ws, r, r, 1, 4)
            r += 1
        # 月小計
        ws.cell(row=r, column=1, value=f"{month}月小計")
        if rows_this_month:
            ws.cell(row=r, column=4, value=f"=SUM(D{start}:D{r - 1})").number_format = AMT_FMT
        style_range(ws, r, r, 1, 4, fill=TOTAL_FILL, font=BOLD)
        subtotal_rows.append(r)
        r += 1

    # 總結
    ws.cell(row=r, column=1, value="總結")
    ws.cell(row=r, column=4, value="=" + "+".join(f"D{x}" for x in subtotal_rows)).number_format = AMT_FMT
    style_range(ws, r, r, 1, 4, fill=TOTAL_FILL, font=BOLD)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 16
    ws.freeze_panes = "A2"
    return ws


# ----------------------------------------------------------------------
INPUT_NAME = "支出.xlsx"
OUTPUT_NAME = "月份彙總.xlsx"


def roc_year(western):
    return western - 1911


def find_year_dir(base, year):
    """找年份資料夾：先找民國年（115），找不到再找西元年（2026）。"""
    for name in (str(roc_year(year)), str(year)):
        d = base / name
        if d.is_dir():
            return d
    return None


def process_year(year, base, strip_month=True, year_filter=None):
    d = find_year_dir(base, year)
    label = f"民國 {roc_year(year)} 年（{year}）"
    if d is None:
        print(f"－ {label}：找不到資料夾 {base / str(roc_year(year))}，略過")
        return False
    src = d / INPUT_NAME
    if not src.is_file():
        print(f"－ {label}：資料夾 {d.name}/ 裡沒有 {INPUT_NAME}，略過")
        return False

    data, products = read_expenses(src, year_filter, strip_month=strip_month)
    if not products:
        print(f"－ {label}：{src} 找不到含「日期」「商品名稱」的工作表，略過")
        return False

    out = d / OUTPUT_NAME
    wb = openpyxl.Workbook()
    write_wide_sheet(wb, data, products)
    write_long_sheet(wb, data, products)
    try:
        wb.save(out)
    except PermissionError:
        print(f"✘ {label}：無法寫入 {out}，請先關閉 Excel 裡開著的檔案")
        return False

    months_with_data = sorted({m for name in products for m in data[name]})
    grand = sum(rec["amt"] for name in products for rec in data[name].values())
    print(f"✔ {label}：{src} → {out}")
    print(f"    產品數：{len(products)}，有資料的月份：{[f'{m}月' for m in months_with_data]}")
    print(f"    總金額：{grand:,.2f}")
    return out


def git_pull(repo_dir):
    """執行前先從 GitHub 拉最新版。失敗（沒網路、有衝突、不是 git 目錄）只警告，不中斷。"""
    if not (repo_dir / ".git").exists():
        print(f"－ {repo_dir} 不是 git 目錄，略過 git pull")
        return
    try:
        r = subprocess.run(["git", "-C", str(repo_dir), "pull", "--ff-only"],
                           capture_output=True, text=True, timeout=60)
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"⚠ git pull 失敗：{e}")
        return
    if r.returncode == 0:
        print("✔ git pull：" + (r.stdout.strip().splitlines() or ["已是最新"])[-1])
    else:
        print("⚠ git pull 失敗，使用本機現有檔案繼續：")
        print("    " + (r.stderr.strip() or r.stdout.strip()).replace("\n", "\n    "))


def git_push(repo_dir, paths, message):
    """把產出的檔案 commit 並推上 GitHub。"""
    rel = [str(p.relative_to(repo_dir)) for p in paths]
    run = lambda *a: subprocess.run(["git", "-C", str(repo_dir), *a], capture_output=True, text=True)
    run("add", "--", *rel)
    if run("diff", "--cached", "--quiet").returncode == 0:
        print("－ 產出檔沒有變更，不需要 push")
        return
    r = run("commit", "-m", message)
    if r.returncode != 0:
        print("⚠ git commit 失敗：", r.stderr.strip()); return
    r = run("push")
    print("✔ 已 push 到 GitHub" if r.returncode == 0 else "⚠ git push 失敗：\n    " + r.stderr.strip())


def main():
    ap = argparse.ArgumentParser(
        description="讀取各年份資料夾裡的 支出.xlsx，在同一個資料夾產出 月份彙總.xlsx。"
                    "預設處理今年與前一年（資料夾名稱為民國年，例如 115、114）。")
    ap.add_argument("years", nargs="*", type=int,
                    help="要處理的年份（民國或西元皆可，例如 114 115 或 2025 2026）；不填則為今年與前一年")
    ap.add_argument("--base", default=None,
                    help="年份資料夾所在的目錄（預設為本程式所在的目錄）")
    ap.add_argument("--year-filter", action="store_true",
                    help="只統計日期落在該年份的資料列（預設不過濾，以日期月份歸類）")
    ap.add_argument("--keep-month-suffix", action="store_true",
                    help="保留產品名稱裡的月份字樣（預設會把「勞保費-1月」合併成「勞保費」）")
    ap.add_argument("--no-pull", action="store_true", help="執行前不要先 git pull")
    ap.add_argument("--push", action="store_true", help="產出後自動 commit 並 push 月份彙總.xlsx")
    args = ap.parse_args()

    script_dir = Path(__file__).resolve().parent
    base = Path(args.base) if args.base else script_dir
    if not args.no_pull:
        git_pull(script_dir)

    if args.years:
        years = [y + 1911 if y < 1000 else y for y in args.years]  # 民國年轉西元
    else:
        this_year = dt.date.today().year
        years = [this_year, this_year - 1]

    outputs = []
    for y in years:
        out = process_year(y, base, strip_month=not args.keep_month_suffix,
                           year_filter=y if args.year_filter else None)
        if out:
            outputs.append(out)
    if not outputs:
        raise SystemExit("沒有產出任何檔案。")
    if args.push:
        git_push(script_dir, outputs, f"Update monthly rollup ({dt.date.today():%Y-%m-%d})")


if __name__ == "__main__":
    main()
